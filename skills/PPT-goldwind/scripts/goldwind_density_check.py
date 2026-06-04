#!/usr/bin/env python3
"""Audit Goldwind native PPTX decks for realistic reporting density.

This check is intentionally heuristic. It catches the failure mode where a deck
is structurally valid but each content slide contains only a few sparse cards or
an empty generated-image frame.
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path

try:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
except ModuleNotFoundError as exc:  # pragma: no cover - environment guard
    raise SystemExit("python-pptx is required. Install with: python3 -m pip install python-pptx") from exc


COPYRIGHT = "GOLDWIND SCIENCE"


def text_of(shape) -> str:
    if not getattr(shape, "has_text_frame", False):
        return ""
    return shape.text.replace("\n", " ").strip()


def is_content_text(text: str) -> bool:
    if not text:
        return False
    upper = text.upper()
    if COPYRIGHT in upper:
        return False
    if text.startswith("Source:"):
        return False
    if text in {"THANKS", "目录"}:
        return False
    return True


def check(path: Path, min_text_shapes: int, min_chars: int, min_image_text_shapes: int) -> list[str]:
    errors: list[str] = []
    prs = Presentation(str(path))
    if len(prs.slides) < 4:
        return [f"Deck has too few slides for density audit: {len(prs.slides)}"]

    for zero_idx in range(2, len(prs.slides) - 1):
        slide_num = zero_idx + 1
        slide = prs.slides[zero_idx]
        text_items = [text_of(shape) for shape in slide.shapes]
        content_texts = [text for text in text_items if is_content_text(text)]
        text_shape_count = len(content_texts)
        char_count = sum(len(text) for text in content_texts)
        picture_count = sum(1 for shape in slide.shapes if shape.shape_type == MSO_SHAPE_TYPE.PICTURE)

        if text_shape_count < min_text_shapes:
            errors.append(
                f"Slide {slide_num}: low text density ({text_shape_count} content text shapes; "
                f"expected >= {min_text_shapes})."
            )
        if char_count < min_chars:
            errors.append(
                f"Slide {slide_num}: low content volume ({char_count} chars; expected >= {min_chars})."
            )
        if picture_count and text_shape_count < min_image_text_shapes:
            errors.append(
                f"Slide {slide_num}: image-bearing page lacks editable overlay/detail text "
                f"({text_shape_count}; expected >= {min_image_text_shapes})."
            )
    return errors


def main() -> int:
    parser = argparse.ArgumentParser(description="Check Goldwind PPTX content density.")
    parser.add_argument("pptx", type=Path)
    parser.add_argument("--min-text-shapes", type=int, default=20)
    parser.add_argument("--min-chars", type=int, default=320)
    parser.add_argument("--min-image-text-shapes", type=int, default=24)
    args = parser.parse_args()

    if not args.pptx.exists():
        print(f"Goldwind density check FAILED: file not found: {args.pptx}", file=sys.stderr)
        return 1
    errors = check(args.pptx, args.min_text_shapes, args.min_chars, args.min_image_text_shapes)
    if errors:
        print("Goldwind density check FAILED", file=sys.stderr)
        for error in errors:
            print(f"- {error}", file=sys.stderr)
        return 1
    print("Goldwind density check PASSED")
    print(f"- file: {args.pptx}")
    print(f"- minimum content text shapes per content slide: {args.min_text_shapes}")
    print(f"- minimum content chars per content slide: {args.min_chars}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
